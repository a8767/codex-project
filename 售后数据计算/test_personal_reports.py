import importlib
import unittest
from pathlib import Path
from unittest.mock import patch
from contextlib import contextmanager
from uuid import uuid4


@contextmanager
def fixture_directory():
    folder = Path(__file__).parent / '个人统计更新核验' / '绩效改造' / ('test_' + uuid4().hex[:10])
    folder.mkdir(parents=True)
    yield folder


class PersonalReportTests(unittest.TestCase):
    def test_attribution_and_no_group_allocation(self):
        for module_name in ('generate_monthly_reports', '影刀_售后月度数据计算'):
            module = importlib.import_module(module_name)
            module.reset_missing_inputs()
            roster = [
                {'售后组': '售后A1组', '组员详情': '张三,李四', '组长': '张三'},
                {'售后组': '售后A2组', '组员详情': '张三，王五', '组长': '赵六'},
            ]
            with patch.object(module, 'read_records', return_value=roster):
                rows = module.personal_report_rows(
                    Path('.'), None,
                    [{'登记人': '张三'}, {'登记人': ''}],
                    [{'登记人': '张三', '店铺': '整车店'},
                     {'登记人': '张三', '店铺': '配件店'},
                     {'登记人': '名单外', '店铺': '整车店'}],
                    [{'登记人': '张三', '补偿金额': 1.25},
                     {'登记人': '张三', '补偿金额': 2.75}],
                    {'配件店'}, '2026-09', '2026-09-22')
            result = {r['姓名']: r for r in rows}
            self.assertEqual(len(result), 6)
            self.assertEqual(result['张三']['售后组'], '售后A1组、售后A2组')
            self.assertEqual(result['张三']['退货登记次数'], 1)
            self.assertEqual(result['张三']['补发登记次数'], 2)
            self.assertEqual(result['张三']['配件店铺补发次数'], 1)
            self.assertEqual(result['张三']['补偿总金额'], 4)
            self.assertEqual(result['张三']['补偿涉及订单数'], 0)
            self.assertIsNone(result['张三']['每单平均补偿金额'])
            self.assertEqual(result['李四']['退货登记次数'], 0)
            self.assertEqual(result['赵六']['售后组'], '售后A2组')
            self.assertEqual(result['未填写登记人']['退货登记次数'], 1)
            self.assertEqual(result['名单外']['补发登记次数'], 1)
            for row in rows:
                for field in ('营业额', '订单量', '万里牛补发', '退货率', '补发率'):
                    self.assertNotIn(field, row)

    def test_order_dedup_aliases_and_threshold(self):
        from decimal import Decimal
        for module_name in ('generate_monthly_reports', '影刀_售后月度数据计算'):
            module = importlib.import_module(module_name)
            module.reset_missing_inputs()
            roster = [{'售后组': '售后A1组', '组员详情': '张三,李四'}]
            common = {'店铺': '甲店', '订单号': '001', '登记时间': '2026-09-21', '原因': '质量问题'}
            comp = [{**common, '登记人': '小张', '补偿金额': 20},
                    {**common, '登记人': '李四', '补偿金额': 30},
                    {**common, '登记人': '张三', '订单号': '', '补偿金额': 100}]
            excluded = {**common, '登记人': '张三', '补偿金额': 80, '是否不参与计算判断': '是'}
            resend = [{**common, '登记人': '小张'}, {**common, '登记人': '张三'},
                      {**common, '店铺': '乙店', '登记人': '张三'}]
            with patch.object(module, 'read_records', return_value=roster), patch.object(
                    module, 'personal_settings', return_value=({'小张': '张三'}, Decimal(50), [])):
                rows = module.personal_report_rows(Path('.'), None, [], resend, comp, set(), '2026-09', '2026-09-21', comp + [excluded])
                again = module.personal_report_rows(Path('.'), None, [], list(reversed(resend)), comp, set(), '2026-09', '2026-09-21', comp + [excluded])
            data = {r['姓名']: r for r in rows}
            self.assertEqual(len(data), 2)
            self.assertEqual(data['张三']['补发登记次数'], 3)
            self.assertEqual(data['张三']['补发涉及订单数'], 2)
            self.assertEqual(data['张三']['重复补发订单数'], 1)
            self.assertEqual(data['张三']['补偿总金额'], 120)
            self.assertEqual(data['张三']['每单平均补偿金额'], 20)
            self.assertEqual(data['张三']['高额补偿订单数'], 1)
            self.assertEqual(data['李四']['高额补偿订单数'], 1)
            self.assertEqual(data['张三']['未计入补偿金额'], 80)
            ids = lambda result: {d['记录ID'] for r in result for d in r['_考核明细']}
            self.assertEqual(ids(rows), ids(again))
            self.assertTrue(any('多人登记' in d['异常提示'] for d in data['李四']['_考核明细']))

    def test_settings_reject_conflicting_aliases(self):
        import tempfile
        from openpyxl import Workbook
        for module_name in ('generate_monthly_reports', '影刀_售后月度数据计算'):
            module = importlib.import_module(module_name)
            with fixture_directory() as folder:
                root = Path(folder)
                w = Workbook()
                s = w.active
                s.title = '考核设置'
                s['B2'] = 'invalid'
                for index, row in enumerate([['小张', '张三', '是'], ['小张', '李四', '是']], 7):
                    for col, value in enumerate(row, 1):
                        s.cell(index, col, value)
                w.save(root / '售后个人数据统计表.xlsx')
                with self.assertRaisesRegex(ValueError, '冲突'):
                    module.personal_settings(root)

    def test_review_survives_reordered_export_and_archives(self):
        import tempfile
        from openpyxl import load_workbook
        module = importlib.import_module('generate_monthly_reports')
        template = Path(__file__).parent / '数据表模板/售后个人数据统计表.xlsx'
        roster = [{'售后组': '售后A1组', '组员详情': '张三'}]
        records = [{'登记人': '张三', '店铺': '甲', '订单号': '1', '原因': '', '登记时间': '2026-09-21'},
                   {'登记人': '张三', '店铺': '甲', '订单号': '2', '原因': '', '登记时间': '2026-09-21'}]
        with patch.object(module, 'read_records', return_value=roster):
            rows = module.personal_report_rows(Path('.'), None, [], records, [], set(), '2026-09', '2026-09-21')
        with fixture_directory() as folder:
            output = Path(folder) / 'result.xlsx'
            module.write_personal_report(template, output, rows)
            w = load_workbook(output)
            s = w['考核异常明细']
            key = s['A2'].value
            s['M2'] = '客服责任'
            s['N2'] = '张三'
            s['O2'] = '已确认'
            s['P2'] = '核实为漏登记'
            w['个人绩效汇总']['AB2'] = '保留评语'
            w.save(output)
            w.close()
            rows[0]['_考核明细'].reverse()
            module.write_personal_report(template, output, rows)
            w = load_workbook(output)
            data = list(w['考核异常明细'].values)
            saved = next(dict(zip(data[0], r)) for r in data[1:] if r[0] == key)
            self.assertEqual(saved['复核备注'], '核实为漏登记')
            self.assertEqual(w['个人绩效汇总']['AB2'].value, '保留评语')
            w.close()
            rows[0]['_考核明细'] = [d for d in rows[0]['_考核明细'] if d['记录ID'] != key]
            module.write_personal_report(template, output, rows)
            w = load_workbook(output)
            self.assertEqual(w['复核留存']['A2'].value, key)
            w.close()


if __name__ == '__main__':
    unittest.main()
