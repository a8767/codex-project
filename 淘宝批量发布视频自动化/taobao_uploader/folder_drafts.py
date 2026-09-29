"""Create reviewable upload drafts from the newest relevant reference data.

The reference database supplies model, ranking, topic, and curated-copy evidence.
Unverified model, price, scene, and specification claims are not added to copy.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from hashlib import sha1
from pathlib import Path
import re
import sqlite3
from time import time_ns

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .copy_library import (
    BikeVideoCopy,
    has_named_copy_theme,
    load_bike_video_copies,
    load_copy_templates,
    suggest_bike_video_copy,
    suggest_copy,
)
from .reference_match import ReferenceAdvice, recommend
from .tasks import HEADERS


MODEL_TOKEN = re.compile(r"(?<![A-Za-z0-9])([A-Za-z]{1,10}[._ -]*\d{1,6}[A-Za-z0-9_-]*)(?![A-Za-z0-9])")
GENERIC_NAMES = {"视频", "素材", "原视频", "待发布", "氛围感", "氛围感视频", "骑行", "骑行视频", "山地车", "自行车"}
CATEGORIES = (
    ("山地车", ("山地车", "山地自行车")),
    ("公路车", ("公路车", "公路自行车")),
    ("折叠车", ("折叠车", "折叠自行车")),
    ("儿童自行车", ("儿童自行车", "童车")),
    ("自行车", ("自行车", "单车")),
)
PREFERRED_TOPIC = {
    "山地车": ("山地车", "山地自行车", "骑行"),
    "公路车": ("公路自行车", "公路车", "骑行"),
    "折叠车": ("折叠自行车", "折叠车", "骑行"),
    "儿童自行车": ("儿童自行车", "骑行"),
    "自行车": ("自行车", "骑行"),
}
EXTRA_HEADERS = ("参考标题", "匹配说明", "榜单参考标题", "榜单参考文案", "榜单参考来源", "上传平台")
# These are complete topic names confirmed in the current Taobao publisher.
# Reference-database hashtags such as “山地车” are not selectable activity topics.
PUBLISHER_TOPICS = {"山地车": "山地骑行的快乐"}
GENERIC_PUBLISHER_TOPIC = "进来一起享受骑行"
CONTENT_TAGS = {
    "山地车": ("山地自行车", "骑行", "户外骑行"),
    "公路车": ("公路自行车", "骑行"),
    "折叠车": ("折叠自行车", "骑行"),
    "儿童自行车": ("儿童自行车", "童车", "骑行"),
    "自行车": ("自行车", "骑行"),
}
GENERIC_CONTENT_TAGS = ("骑行", "自行车")


@dataclass(frozen=True)
class DraftResult:
    output: Path
    total: int
    model_matched: int
    generic: int
    no_reference: int


@dataclass(frozen=True)
class ModelMatch:
    model: str
    category: str
    topic: str
    reference_count: int
    reference_title: str
    related_topics: tuple[str, ...]


def _model_pattern(model: str) -> re.Pattern[str]:
    chunks = re.findall(r"[A-Za-z]+|\d+", model)
    return re.compile(r"(?<![A-Za-z0-9])" + r"[._ -]*".join(map(re.escape, chunks)) +
                      r"(?![A-Za-z0-9])", re.IGNORECASE)


def _catalog_model_key(name: str, connection: sqlite3.Connection) -> str | None:
    exists = connection.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='model_catalog'"
    ).fetchone()
    if not exists:
        return None
    columns = {row[1] for row in connection.execute("PRAGMA table_info(model_catalog)")}
    searchable = [column for column in ("model_key", "model_name", "chinese_code", "english_code")
                  if column in columns]
    if not searchable:
        return None
    condition = " OR ".join(f'"{column}"=?' for column in searchable)
    row = connection.execute(
        f"SELECT model_key FROM model_catalog WHERE {condition} LIMIT 1",
        tuple(name for _ in searchable),
    ).fetchone()
    if row:
        return row[0] or name
    # Some asset folders use a short product nickname, such as "绯影";
    # accept it only when it resolves to one catalog model key.
    if re.search(r"[\u3400-\u9fff]", name) and "model_name" in columns:
        matches = {row[0] for row in connection.execute(
            "SELECT DISTINCT model_key FROM model_catalog WHERE model_name LIKE ?",
            (f"%{name}%",),
        ) if row[0]}
        if len(matches) == 1:
            return matches.pop()
    return None


def _folder_model(name: str, connection: sqlite3.Connection) -> str | None:
    token = MODEL_TOKEN.search(name)
    if token:
        raw = token.group(1).strip("._ -")
        candidates = (raw, re.sub(r"[-_ ]+\d{1,4}$", "", raw).strip("._ -"))
        for candidate in candidates:
            model_key = _catalog_model_key(candidate, connection)
            if model_key:
                return model_key
        return raw
    cleaned = re.sub(r"(?:视频|素材|片段|成片)$", "", name).strip(" _-（）()")
    if cleaned in GENERIC_NAMES or len(cleaned) < 2 or len(cleaned) > 20:
        return None
    # Chinese model names count only when they actually occur in the reference bank.
    if re.search(r"[\u3400-\u9fff]", cleaned):
        candidates = (cleaned, re.sub(r"[-_ ]+\d{1,4}$", "", cleaned).strip("._ -"))
        for candidate in candidates:
            model_key = _catalog_model_key(candidate, connection)
            if model_key:
                return model_key
        count = connection.execute("SELECT count(*) FROM videos WHERE title LIKE ?", (f"%{cleaned}%",)).fetchone()[0]
        if count >= 2:
            return cleaned
    return None


def _filename_model(video: Path, connection: sqlite3.Connection) -> str | None:
    """Recognize a model prefix from names such as GX1060-12.mp4."""
    numbered_clip = re.match(r"^(.+?)[\s_-]+\d{1,4}$", video.stem)
    if not numbered_clip:
        return None
    candidate = numbered_clip.group(1).strip("._ -")
    if not candidate or candidate.isdigit():
        return None
    if re.search(r"[A-Za-z]", candidate):
        model_key = _catalog_model_key(candidate, connection)
        if model_key:
            return model_key
    return _folder_model(candidate, connection)


def _find_model_folder(video: Path, root: Path, connection: sqlite3.Connection) -> str | None:
    try:
        relative = video.relative_to(root)
        names = [part for part in relative.parts[:-1]][::-1] + [root.name]
    except ValueError:
        # A multi-select dialog can choose files from separate folders/drives.
        # Inspect their ancestor names independently in that case.
        names = list(video.parent.parts[::-1])
    for name in names:
        model = _folder_model(name, connection)
        if model:
            return model
    return _filename_model(video, connection)


def _match_model(connection: sqlite3.Connection, model: str,
                 target_platform: str = "淘宝光合") -> ModelMatch:
    video_columns = {row[1] for row in connection.execute("PRAGMA table_info(videos)")}
    platform_expression = "platform" if "platform" in video_columns else "''"
    if re.search(r"[A-Za-z]", model):
        prefix = re.match(r"[A-Za-z]+", model)
        candidates = connection.execute(
            f"SELECT video_key,title,{platform_expression} FROM videos WHERE title LIKE ?",
                                        (f"%{prefix.group(0)}%",)).fetchall()
        pattern = _model_pattern(model)
        matching = [(key, title, platform) for key, title, platform in candidates
                    if pattern.search(title)]
    else:
        matching = connection.execute(
            f"SELECT video_key,title,{platform_expression} FROM videos WHERE title LIKE ?",
            (f"%{model}%",),
        ).fetchall()
    platform_refs = [(key, title) for key, title, platform in matching
                     if platform == target_platform]
    preferred_source_types = {
        "淘宝光合": ("淘宝优质视频",),
        "抖音": ("热门视频榜", "看后搜视频榜", "视频销量榜"),
    }.get(target_platform, ())
    preferred_keys: set[str] = set()
    if preferred_source_types:
        if connection.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='video_sources'"
        ).fetchone():
            placeholders = ",".join("?" for _ in preferred_source_types)
            preferred_keys.update(key for (key,) in connection.execute(
                f"SELECT DISTINCT video_key FROM video_sources WHERE source_type IN ({placeholders})",
                preferred_source_types,
            ))
        if target_platform == "抖音" and connection.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='rank_occurrences'"
        ).fetchone():
            preferred_keys.update(key for (key,) in connection.execute(
                "SELECT DISTINCT video_key FROM rank_occurrences"
            ))
    preferred_refs = [(key, title) for key, title in platform_refs if key in preferred_keys]
    # Prefer the selected platform's own model examples; cross-platform examples
    # are only a fallback when its quality/ranking set has no matching model.
    refs = preferred_refs or platform_refs or [(key, title) for key, title, _ in matching]
    if not refs:
        return ModelMatch(model, "", "骑行", 0, "", ())

    category_counts = Counter()
    for _, title in refs:
        for category, words in CATEGORIES:
            if any(word in title for word in words):
                category_counts[category] += 1
    specific = [name for name, _ in CATEGORIES if name != "自行车"]
    category = max(specific, key=lambda x: category_counts[x]) if specific else ""
    if category_counts[category] < max(2, len(refs) // 5):
        category = "自行车" if category_counts["自行车"] else ""
    if not category and connection.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='store_video_metrics'"
    ).fetchone():
        metric_columns = {row[1] for row in connection.execute("PRAGMA table_info(store_video_metrics)")}
        if {"product_name", "statistics_period"}.issubset(metric_columns):
            latest_period = connection.execute(
                "SELECT MAX(statistics_period) FROM store_video_metrics"
            ).fetchone()[0]
            if latest_period:
                pattern = _model_pattern(model)
                names = [row[0] for row in connection.execute(
                    "SELECT DISTINCT product_name FROM store_video_metrics "
                    "WHERE statistics_period=? AND product_name LIKE ?",
                    (latest_period, f"%{model}%"),
                ) if pattern.search(row[0] or "")]
                metric_categories = Counter(
                    kind for name in names
                    if (kind := _category(name or "", "")) in {*specific, "自行车"}
                )
                if metric_categories:
                    category = metric_categories.most_common(1)[0][0]

    keys = [key for key, _ in refs]
    topic_counts: Counter[str] = Counter()
    for start in range(0, len(keys), 500):
        part = keys[start:start + 500]
        placeholders = ",".join("?" for _ in part)
        topic_counts.update(topic for (topic,) in connection.execute(
            f"SELECT topic FROM video_topics WHERE video_key IN ({placeholders})", part))
    preferred = PREFERRED_TOPIC.get(category, ("骑行", "自行车"))
    topic = next((candidate for candidate in preferred if topic_counts[candidate]), "骑行")
    related = tuple(topic for topic, _ in topic_counts.most_common(8)
                    if topic not in {"直播", "直播间热卖中"})[:4]

    # Reference title is evidence for the human reviewer, never an auto-published copy.
    published_at = {}
    if "published_at" in video_columns and refs:
        placeholders = ",".join("?" for _ in refs)
        published_at = dict(connection.execute(
            f"SELECT video_key,published_at FROM videos WHERE video_key IN ({placeholders})",
            [key for key, _ in refs],
        ).fetchall())
    newest_date = max((value or "" for value in published_at.values()), default="")
    newest_refs = [(key, title) for key, title in refs
                   if (published_at.get(key) or "") == newest_date] if newest_date else refs
    clean_titles = [re.sub(r"\s*[@#＃].*$", "", title).strip(" -，。！! ")
                    for _, title in newest_refs]
    short = [title for title in clean_titles if 4 <= len(title) <= 30]
    reference_title = min(short or clean_titles, key=lambda x: (len(x), x))
    return ModelMatch(model, category, topic, len(refs), reference_title, related)


def _select_copy(templates, bike_copies: tuple[BikeVideoCopy, ...], category: str,
                 model: str | None, folder_name: str, video_name: str, index: int,
                 themes, used_copy_ids: set[str], rotation_key: str):
    selected = suggest_bike_video_copy(
        bike_copies, category, folder_name, video_name, index, themes,
        used_copy_ids, rotation_key,
    )
    if selected:
        used_copy_ids.add(selected.copy_id)
        return (selected.title, selected.body, selected.direction,
                selected.video_theme, selected.usage_scenario)
    title, body = suggest_copy(templates, category, model, index, themes)
    direction = ("带货种草向" if index % 2 else "氛围感向") if not model else "通用草稿模板"
    return title, body, direction, "通用（文件名未识别主题）", "使用前请核对视频内容及商品信息"


def _copy_title_candidates(title: str, body: str) -> tuple[str, ...]:
    """Use a short body line as a title when it helps avoid title repetition."""
    candidates = [title.strip()[:30]]
    exact_body = body.strip()
    if 0 < len(exact_body) <= 30:
        candidates.append(exact_body)
    else:
        opening = re.split(r"[。！!？?；;\n]", exact_body, maxsplit=1)[0].strip(" ，,：:")
        if 0 < len(opening) <= 30:
            candidates.append(opening)
    return tuple(dict.fromkeys(value for value in candidates if value))


def _topic_caption(content_tags: tuple[str, ...]) -> str:
    """Format confirmed content tags as a concise caption when copy becomes title."""
    tags = []
    for value in content_tags:
        tag = value.strip().lstrip("#＃")
        if tag and tag not in tags:
            tags.append("#" + tag)
    return " ".join(tags)


def generate_folder_draft(folder: Path, reference_db: Path, output: Path,
                          target_platform: str = "淘宝光合",
                          video_paths: list[Path] | None = None) -> DraftResult:
    folder, reference_db, output = folder.resolve(), reference_db.resolve(), output.resolve()
    if not folder.is_dir():
        raise ValueError(f"视频文件夹不存在：{folder}")
    if not reference_db.is_file():
        raise ValueError(f"找不到视频发布参考数据库：{reference_db}")
    if video_paths is None:
        videos = [p.resolve() for p in folder.rglob("*")
                  if p.is_file() and p.suffix.lower() == ".mp4"]
    else:
        videos = list(dict.fromkeys(Path(p).expanduser().resolve() for p in video_paths))
        invalid = [str(p) for p in videos if not p.is_file() or p.suffix.lower() != ".mp4"]
        if invalid:
            raise ValueError("所选文件中包含不存在的文件或非 MP4 视频：" + "、".join(invalid))
    videos.sort(key=lambda p: str(p).casefold())
    if not videos:
        raise ValueError(f"文件夹内没有 MP4 视频：{folder}")
    if output == reference_db:
        raise ValueError("草稿输出路径不能覆盖参考数据库")
    output.parent.mkdir(parents=True, exist_ok=True)

    connection = sqlite3.connect(f"file:{reference_db.as_posix()}?mode=ro", uri=True)
    try:
        copy_templates = load_copy_templates(connection)
        bike_copies = load_bike_video_copies(connection)
        used_copy_ids: set[str] = set()
        # Vary near-equal candidates between separate draft generations while
        # keeping a single generation internally consistent.
        rotation_key = sha1(f"{time_ns()}:{folder}:{output}".encode("utf-8")).hexdigest()
        matches: dict[str, ModelMatch] = {}
        advice_by_group: dict[str, ReferenceAdvice] = {}
        sequence: defaultdict[str, int] = defaultdict(int)
        copy_sequence = 0
        used_titles: set[str] = set()
        used_bodies: set[str] = set()
        counts = Counter()
        records = []
        for video in videos:
            copy_sequence += 1
            model = _find_model_folder(video, folder, connection)
            folder_name = video.parent.name
            group = model or f"generic:{folder_name}"
            sequence[group] += 1
            if model:
                match = matches.get(model)
                if match is None:
                    match = _match_model(connection, model, target_platform)
                    matches[model] = match
                counts["model"] += 1
                if not match.reference_count:
                    counts["no_reference"] += 1
                advice = advice_by_group.get(group)
                if advice is None:
                    advice = recommend(connection, model, match.category, folder_name,
                                       target_platform=target_platform)
                    advice_by_group[group] = advice
                topic = (advice.topic if advice.count else
                         PUBLISHER_TOPICS.get(match.category, GENERIC_PUBLISHER_TOPIC))
                content_tags = (*CONTENT_TAGS.get(match.category, GENERIC_CONTENT_TAGS), f"#{model}")
                note = (f"车型={model}；参考视频={match.reference_count}；榜单话题=" +
                        "、".join(advice.related_topics or match.related_topics or (match.topic,)) +
                        f"；平台/榜单同类参考={advice.count}；来源={advice.source or '无'}；"
                        f"主题={advice.themes[0][0] if advice.themes else '通用'}；"
                        f"平台话题={topic}；内容标签={'、'.join(content_tags)}；车型号以 #{model} 标签填写；请核对视频内容")
                reference_title = match.reference_title
            else:
                counts["generic"] += 1
                advice = advice_by_group.get(group)
                if advice is None:
                    advice = recommend(connection, None, "", folder_name,
                                       target_platform=target_platform)
                    advice_by_group[group] = advice
                topic = advice.topic
                content_tags = (*GENERIC_CONTENT_TAGS, f"#{model}") if model else GENERIC_CONTENT_TAGS
                note = ("文件夹未识别明确车型；不补写未经视频或商品信息确认的车型、配置或价格；"
                        f"平台/榜单通用参考={advice.count}；来源={advice.source or '无'}；"
                        f"主题={advice.themes[0][0] if advice.themes else '通用'}；"
                        f"榜单话题={'、'.join(advice.related_topics) or '无'}；"
                        f"平台话题={topic}；内容标签={'、'.join(content_tags)}；"
                        f"车型号标签={'#' + model if model else '无'}；请核对视频内容")
                reference_title = ""
            copy_category = match.category if model else ""
            # Ranking themes describe reference videos, not these unseen files.
            # Let filename scenes drive copy-theme selection; otherwise keep a
            # neutral template and expose the missing theme in the match note.
            copy_themes = advice.themes if model and has_named_copy_theme(folder_name, video.stem) else ()
            title, body, copy_direction, copy_theme, copy_scenario = _select_copy(
                copy_templates, bike_copies, copy_category, model, folder_name,
                video.stem, copy_sequence, copy_themes, used_copy_ids, rotation_key,
            )
            note += (f"；原创文案方向={copy_direction}；适用文案主题={copy_theme}"
                     f"；使用场景={copy_scenario}")
            repeated_title_fallback = None
            for _ in range(max(100, len(videos) * 4)):
                if body not in used_bodies:
                    title_options = _copy_title_candidates(title, body)
                    distinct_title = next((candidate for candidate in title_options
                                           if candidate not in used_titles), None)
                    if distinct_title:
                        title = distinct_title
                        break
                    if repeated_title_fallback is None:
                        repeated_title_fallback = (
                            title_options[0], body, copy_direction, copy_theme, copy_scenario
                        )
                sequence[group] += 1
                copy_sequence += 1
                title, body, copy_direction, copy_theme, copy_scenario = _select_copy(
                    copy_templates, bike_copies, copy_category, model, folder_name,
                    video.stem, copy_sequence, copy_themes, used_copy_ids, rotation_key,
                )
                note = re.sub(r"；原创文案方向=.*$", "", note)
                note += (f"；原创文案方向={copy_direction}；适用文案主题={copy_theme}"
                         f"；使用场景={copy_scenario}")
            else:
                if repeated_title_fallback is None:
                    raise ValueError(f"文案库不足以为所有视频生成不同草稿：{folder_name}")
                title, body, copy_direction, copy_theme, copy_scenario = repeated_title_fallback
                note = re.sub(r"；原创文案方向=.*$", "", note)
                note += (f"；原创文案方向={copy_direction}；适用文案主题={copy_theme}"
                         f"；使用场景={copy_scenario}")
            used_titles.add(title)
            used_bodies.add(body)
            published_body = (_topic_caption(content_tags) if title.strip() in body.strip()
                              else body)
            try:
                relative = str(video.relative_to(folder)).replace("\\", "/")
            except ValueError:
                relative = str(video).replace("\\", "/")
            task_id = "folder-" + sha1(relative.casefold().encode("utf-8")).hexdigest()[:12]
            records.append((task_id, str(video), title, published_body, "", "", "", ";".join(content_tags), topic, "", "",
                            reference_title, note, advice.title, advice.body, advice.source, target_platform))
    finally:
        connection.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "发布任务"
    ws.append((*HEADERS, *EXTRA_HEADERS))
    for record in records:
        ws.append(record)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:Q{len(records) + 1}"
    widths = (23, 65, 36, 32, 23, 32, 18, 18, 18, 22, 22, 45, 90, 48, 75, 26, 20)
    for column, width in zip(ws.columns, widths):
        ws.column_dimensions[column[0].column_letter].width = width
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="273D5B")
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column in {3, 12, 13, 14, 15})
        ws.row_dimensions[row[0].row].height = 34

    guide = wb.create_sheet("生成说明")
    for line in (
        ("项目", "说明"),
        ("生成依据", f"视频文件夹：{folder}"),
        ("参考数据库", str(reference_db)),
        ("上传平台", target_platform),
        ("匹配规则", "优先识别当前或上级文件夹的车型名；先匹配上传平台自己的车型视频与优质榜单数据；没有足够同平台样本时才回退到其他平台。标题、正文不写车型号；识别到的型号写成独立 #内容标签。"),
        ("文案选择", "发布候选读取自行车视频文案表的24条文案和原创文案库的48条文案；新采集的光合发布视频文案与话题用于同平台主题、话题匹配，并按相对发布时间提高近90天样本权重，不直接复制其中的车型、配置或价格表述。文件名主题明确时优先匹配数据库主题和关键词；文件名无方向或主题时方向交替，并使用中性通用文案。"),
        ("数据时效", "榜单参考优先使用最新采集批次、最新月份记录和最新发布视频；主题及文案候选结合视频话题、关键词和适用场景。"),
        ("关键信息", "新文案不自动填具体车型、价格或配置；只使用视频文件名、目录和已确认商品信息能支持的内容。"),
        ("榜单参考", "榜单参考标题、文案和来源仅供核对；不会把他人的文案或未经证实的配置直接填入发布栏。"),
        ("通用视频", "没有明确车型时使用通用骑行标题，不写车型、配置或价格。"),
        ("内容标签", "按车型类别写入平台可搜索的内容标签；无明确车型时仅使用骑行、自行车。导入前请按画面核对。"),
        ("发布前必须填写", "创作者声明；按视频实际内容核对标题、话题、商品 ID、封面和发布时间。"),
        ("安全默认", "标题、正文是原创草稿文案，须按视频画面核对；商品 ID、创作者声明和发布时间均留空。生成草稿不会上传或发布视频。"),
        ("导入方法", "保存并关闭本工作簿后，在主程序中点击“选择并导入”。"),
    ):
        guide.append(line)
    guide.column_dimensions["A"].width = 23
    guide.column_dimensions["B"].width = 110
    for cell in guide[1]:
        cell.fill = PatternFill("solid", fgColor="273D5B")
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
    wb.save(output)
    return DraftResult(output, len(records), counts["model"], counts["generic"], counts["no_reference"])
