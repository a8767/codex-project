"""Import bicycle model code exports into the video reference database."""

from __future__ import annotations

from pathlib import Path
import re
import sqlite3

from openpyxl import load_workbook


REFERENCE_DB = (Path(__file__).resolve().parents[1] / "outputs" /
                "video_reference_db" / "视频发布参考数据库.sqlite3")


def normalize_model(value: str) -> str:
    return re.sub(r"[\s._-]+", "", value).upper()


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _model_from_code(chinese_code: str) -> str:
    # Some exports have no dedicated model column; their first code token is the model.
    return chinese_code.split(maxsplit=1)[0] if chinese_code else ""


def import_model_workbooks(database: Path, files: list[Path]) -> dict[str, int]:
    """Replace rows from each source workbook atomically; retain sheet and row lineage."""
    database = database.resolve()
    if not database.is_file():
        raise FileNotFoundError(f"找不到视频参考数据库：{database}")
    connection = sqlite3.connect(database)
    counts: dict[str, int] = {}
    try:
        connection.execute("""CREATE TABLE IF NOT EXISTS model_catalog (
            brand TEXT NOT NULL,
            model_name TEXT NOT NULL,
            model_key TEXT NOT NULL,
            chinese_code TEXT NOT NULL,
            english_code TEXT NOT NULL,
            source_file TEXT NOT NULL,
            source_sheet TEXT NOT NULL,
            source_row INTEGER NOT NULL,
            PRIMARY KEY (source_file, source_sheet, source_row)
        )""")
        connection.execute("CREATE INDEX IF NOT EXISTS idx_model_catalog_key ON model_catalog(brand,model_key)")
        for file in files:
            source = file.resolve()
            if not source.is_file():
                raise FileNotFoundError(f"找不到车型表：{source}")
            brand = source.name.split("车型编码", 1)[0]
            if brand not in {"大众", "菲利普", "玛莎拉蒂", "永久"}:
                raise ValueError(f"无法确定车型表品牌：{source.name}")
            workbook = load_workbook(source, read_only=True, data_only=True)
            rows_to_write: list[tuple] = []
            try:
                for sheet in workbook:
                    rows = sheet.iter_rows(values_only=True)
                    headers = tuple(_text(value) for value in next(rows, ()))
                    if "中文编码" not in headers or "英文编码" not in headers:
                        raise ValueError(f"{source.name}/{sheet.title} 缺少中文编码或英文编码列")
                    chinese_index = headers.index("中文编码")
                    english_index = headers.index("英文编码")
                    model_index = next((index for index, title in enumerate(headers)
                                        if title in {"车型", "型号"}), None)
                    for row_number, raw in enumerate(rows, 2):
                        values = [_text(value) for value in raw]
                        chinese = values[chinese_index] if chinese_index < len(values) else ""
                        english = values[english_index] if english_index < len(values) else ""
                        if not chinese and not english:
                            continue
                        model = (values[model_index] if model_index is not None and
                                 model_index < len(values) else "") or _model_from_code(chinese)
                        rows_to_write.append((brand, model, normalize_model(model), chinese,
                                              english, str(source), sheet.title, row_number))
            finally:
                workbook.close()
            with connection:
                connection.execute("DELETE FROM model_catalog WHERE source_file=?", (str(source),))
                connection.executemany("INSERT INTO model_catalog VALUES (?,?,?,?,?,?,?,?)", rows_to_write)
            counts[brand] = len(rows_to_write)
        return counts
    finally:
        connection.close()


def find_model(database: Path, brand: str, model_name: str) -> list[tuple[str, str, str]]:
    """Return exact normalized model matches as (model, Chinese code, English code)."""
    connection = sqlite3.connect(database)
    try:
        return connection.execute("""SELECT DISTINCT model_name,chinese_code,english_code
            FROM model_catalog WHERE brand=? AND model_key=? ORDER BY chinese_code,english_code""",
            (brand, normalize_model(model_name))).fetchall()
    finally:
        connection.close()
