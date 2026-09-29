"""Durable publication ledger and a readable Excel view of confirmed works."""

from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import os
from pathlib import Path
import re
import sqlite3
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .tasks import TZ, Task


DEFAULT_STATS_DIR = Path(__file__).resolve().parents[1] / "outputs" / "publish_stats"
HEADERS = (
    "视频作品文件名称", "作品ID", "作品文案", "作品话题", "视频发布时间（北京时间）", "备注",
)


def _beijing(value: datetime | None) -> str:
    if value is None:
        return ""
    if value.tzinfo is None:
        raise ValueError("统计记录的时间必须包含时区")
    return value.astimezone(TZ).isoformat(timespec="seconds")


def _excel_datetime(value: str) -> datetime | None:
    return datetime.fromisoformat(value).replace(tzinfo=None) if value else None


class PublicationStats:
    def __init__(self, directory: Path = DEFAULT_STATS_DIR):
        self.directory = directory.resolve()
        self.directory.mkdir(parents=True, exist_ok=True)
        self.workbook = self.directory / "发布视频统计表.xlsx"
        self.db = sqlite3.connect(self.directory / "发布视频统计.sqlite3")
        self.db.execute("""CREATE TABLE IF NOT EXISTS publications (
            work_id TEXT PRIMARY KEY,
            source_excel TEXT NOT NULL,
            task_id TEXT NOT NULL,
            title TEXT NOT NULL,
            topic TEXT NOT NULL,
            publish_mode TEXT NOT NULL,
            scheduled_at TEXT NOT NULL,
            submitted_at TEXT NOT NULL,
            verified_at TEXT NOT NULL,
            audit_status TEXT NOT NULL,
            body TEXT NOT NULL,
            product_ids TEXT NOT NULL,
            video_path TEXT NOT NULL,
            cover_path TEXT NOT NULL,
            brand_tags TEXT NOT NULL,
            content_tags TEXT NOT NULL,
            declaration TEXT NOT NULL,
            remark TEXT NOT NULL DEFAULT '',
            UNIQUE(source_excel, task_id)
        )""")
        columns = {row[1] for row in self.db.execute("PRAGMA table_info(publications)")}
        if "remark" not in columns:
            self.db.execute("ALTER TABLE publications ADD COLUMN remark TEXT NOT NULL DEFAULT ''")
        self.db.commit()

    def record(self, task: Task, work_id: str, *, source_excel: Path | None,
               submitted_at: datetime | None, verified_at: datetime,
               audit_status: str = "") -> None:
        if not work_id.isdigit():
            raise ValueError("后台作品 ID 必须是数字字符串")
        source = str(source_excel.resolve()) if source_excel else ""
        previous_task = self.db.execute(
            "SELECT work_id FROM publications WHERE source_excel=? AND task_id=?",
            (source, task.task_id)).fetchone()
        recorded_task_id = task.task_id
        if previous_task and previous_task[0] != work_id:
            existing_ids = [row[0] for row in self.db.execute(
                "SELECT task_id FROM publications WHERE source_excel=?", (source,)
            )]
            attempt_count = sum(
                value == task.task_id or value.startswith(task.task_id + "#attempt-")
                for value in existing_ids
            )
            recorded_task_id = f"{task.task_id}#attempt-{attempt_count + 1}"
        previous_work = self.db.execute(
            "SELECT source_excel,task_id FROM publications WHERE work_id=?", (work_id,)).fetchone()
        if previous_work and previous_work != (source, task.task_id):
            raise ValueError(f"作品 ID {work_id} 已对应另一任务，统计表需要人工核查")
        row = (
            work_id, source, recorded_task_id, task.title, task.topic,
            "定时" if task.publish_at else "立即", _beijing(task.publish_at),
            _beijing(submitted_at), _beijing(verified_at), audit_status,
            task.body, "；".join(task.product_ids), str(task.video),
            str(task.cover) if task.cover else "", "；".join(task.brand_tags),
            "；".join(task.content_tags), task.declaration, "",
        )
        self.db.execute("""INSERT INTO publications VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(work_id) DO UPDATE SET
                verified_at=excluded.verified_at,
                audit_status=CASE WHEN excluded.audit_status='' THEN audit_status
                                  ELSE excluded.audit_status END""", row)
        self.db.commit()

    def record_skipped(self, task: Task, *, source_excel: Path | None,
                       skipped_at: datetime, remark: str = "视频重复上传") -> None:
        """Record a platform duplicate warning without inventing a work ID."""
        source = str(source_excel.resolve()) if source_excel else ""
        key = hashlib.sha256(f"{source}\0{task.task_id}".encode("utf-8")).hexdigest()
        work_id = f"SKIP-{key}"
        previous = self.db.execute(
            "SELECT work_id FROM publications WHERE source_excel=? AND task_id=?",
            (source, task.task_id),
        ).fetchone()
        if previous and previous[0] != work_id:
            raise ValueError(f"任务 {task.task_id} 已有正式作品记录，不能改记为跳过")
        row = (
            work_id, source, task.task_id, task.title, task.topic, "跳过", "", "",
            _beijing(skipped_at), "", task.body, "；".join(task.product_ids),
            str(task.video), str(task.cover) if task.cover else "",
            "；".join(task.brand_tags), "；".join(task.content_tags), task.declaration, remark,
        )
        self.db.execute("""INSERT INTO publications VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(work_id) DO UPDATE SET
                verified_at=excluded.verified_at, remark=excluded.remark""", row)
        self.db.commit()

    def backfill_successes(self, tasks: list[Task], state_db: Path,
                           source_excel: Path) -> int:
        """Recover confirmed works from an older task state without resubmitting."""
        if not state_db.is_file():
            return 0
        source = str(source_excel.resolve())
        imported = 0
        task_db = sqlite3.connect(state_db)
        try:
            for task in tasks:
                row = task_db.execute(
                    "SELECT status,message,updated_at FROM tasks WHERE task_id=?",
                    (task.task_id,)).fetchone()
                if not row or row[0] != "成功":
                    continue
                match = re.search(r"作品 ID\s*(\d+)", row[1])
                if not match:
                    continue
                work_id = match.group(1)
                if self.db.execute(
                        "SELECT 1 FROM publications WHERE work_id=? OR (source_excel=? AND task_id=?)",
                        (work_id, source, task.task_id)).fetchone():
                    continue
                verified_at = datetime.strptime(row[2], "%Y-%m-%d %H:%M:%S").replace(
                    tzinfo=timezone.utc)
                self.record(task, work_id, source_excel=source_excel,
                            submitted_at=None, verified_at=verified_at)
                imported += 1
        finally:
            task_db.close()
        return imported

    def rows(self) -> list[tuple]:
        return self.db.execute("""SELECT video_path,
            CASE WHEN remark<>'' THEN '' ELSE work_id END,
            title,body,topic,content_tags,scheduled_at,submitted_at,verified_at,remark
            FROM publications
            ORDER BY COALESCE(NULLIF(scheduled_at,''),NULLIF(submitted_at,''),verified_at),work_id""").fetchall()

    def export(self) -> Path:
        records = self.rows()
        wb = Workbook()
        ws = wb.active
        ws.title = "发布记录"
        ws.append(HEADERS)
        for record in records:
            video_path, work_id, title, body, topic, content_tags, scheduled, submitted, verified, remark = record
            copy = title if not body or body == title else f"{title}\n{body}"
            if content_tags:
                copy += "\n" + " ".join(f"#{tag.lstrip('#＃').strip()}" for tag in content_tags.split("；")
                                         if tag.lstrip("#＃").strip())
            publish_time = None if remark else _excel_datetime(scheduled or submitted or verified)
            ws.append((Path(video_path).name, work_id, copy, topic, publish_time, remark))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:F{len(records) + 1}"
        ws.sheet_view.showGridLines = False
        widths = (36, 24, 56, 30, 26, 24)
        for column, width in zip(ws.columns, widths):
            ws.column_dimensions[column[0].column_letter].width = width
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="273D5B")
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28
        for cells in ws.iter_rows(min_row=2):
            ws.row_dimensions[cells[0].row].height = 42 if "\n" in cells[2].value else 25
            cells[1].number_format = "@"  # Long platform IDs must remain exact text.
            cells[4].number_format = "yyyy-mm-dd hh:mm"
            for cell in cells:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        with NamedTemporaryFile(prefix="publish-stats-", suffix=".xlsx",
                                dir=self.directory, delete=False) as temporary:
            temporary_path = Path(temporary.name)
        try:
            wb.save(temporary_path)
            os.replace(temporary_path, self.workbook)
        finally:
            wb.close()
            try:
                temporary_path.unlink()
            except FileNotFoundError:
                # Path.unlink(missing_ok=...) is unavailable in older
                # Python runtimes bundled with some ShadowBot installations.
                pass
        return self.workbook

    def close(self) -> None:
        self.db.close()
