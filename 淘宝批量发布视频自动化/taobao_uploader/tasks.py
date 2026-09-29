from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from hashlib import sha256
import json
from pathlib import Path
import re

from openpyxl import load_workbook


HEADERS = (
    "任务编号", "视频路径", "标题", "正文", "商品ID", "封面路径", "品牌标签",
    "内容标签", "话题", "发布时间", "创作者声明",
)
DECLARATIONS = {
    "内容无需标注", "含AI生成内容", "含虚构演绎内容", "内容为转载",
    "个人观点，仅供参考", "内容含营销信息",
}
TZ = timezone(timedelta(hours=8), "Beijing")
TIME_FORMAT = "%Y-%m-%d %H:%M"
MAX_VIDEO_BYTES = 1_500_000_000
SCHEDULE_MIN_LEAD = timedelta(hours=2)
SCHEDULE_MAX_AHEAD = timedelta(days=30)


@dataclass(frozen=True)
class Task:
    row: int
    task_id: str
    video: Path
    title: str
    body: str
    product_ids: tuple[str, ...]
    cover: Path | None
    brand_tags: tuple[str, ...]
    content_tags: tuple[str, ...]
    topic: str
    publish_at: datetime | None
    declaration: str
    upload_platform: str = "淘宝光合"

    @property
    def fingerprint(self) -> str:
        content = {
            "video": str(self.video), "size": self.video.stat().st_size,
            "mtime_ns": self.video.stat().st_mtime_ns, "title": self.title,
            "body": self.body, "products": self.product_ids,
            "cover": str(self.cover) if self.cover else "",
            "brand_tags": self.brand_tags, "content_tags": self.content_tags,
            "topic": self.topic,
            "publish_at": self.publish_at.isoformat() if self.publish_at else "",
            "declaration": self.declaration,
        }
        # Preserve fingerprints of legacy workbooks while isolating future
        # cross-platform task state from otherwise identical Taobao tasks.
        if self.upload_platform != "淘宝光合":
            content["upload_platform"] = self.upload_platform
        return sha256(json.dumps(content, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


class TaskValidationError(ValueError):
    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("\n".join(errors))


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _split(value: str) -> tuple[str, ...]:
    return tuple(part.strip() for part in re.split(r"[;；]", value) if part.strip())


def _path(base: Path, value: str) -> Path:
    path = Path(value).expanduser()
    return (path if path.is_absolute() else base / path).resolve()


def read_tasks(workbook: Path, *, now: datetime | None = None) -> list[Task]:
    workbook = workbook.resolve()
    if not workbook.is_file():
        raise TaskValidationError([f"找不到 Excel：{workbook}"])
    wb = load_workbook(workbook, read_only=True, data_only=True)
    try:
        sheet = wb.active
        rows = sheet.iter_rows(values_only=True)
        first = next(rows, None)
        if not first:
            raise TaskValidationError(["Excel 为空"])
        actual = tuple(_text(x) for x in first[:len(HEADERS)])
        if actual != HEADERS:
            raise TaskValidationError(["第一行表头应依次为：" + "、".join(HEADERS)])
        full_header = tuple(_text(x) for x in first)
        platform_index = full_header.index("上传平台") if "上传平台" in full_header else None
        errors: list[str] = []
        tasks: list[Task] = []
        ids: set[str] = set()
        current = now or datetime.now(TZ)
        if current.tzinfo is None:
            current = current.replace(tzinfo=TZ)
        for row_number, raw in enumerate(rows, 2):
            values = tuple(_text(x) for x in raw[:len(HEADERS)])
            if not any(values):
                continue
            values += ("",) * (len(HEADERS) - len(values))
            task_id, video_s, title, body, products_s, cover_s, brands_s, tags_s, topic, time_s, declaration = values
            upload_platform = (
                _text(raw[platform_index]) if platform_index is not None and platform_index < len(raw)
                else "淘宝光合"
            )
            row_errors: list[str] = []
            if not task_id:
                row_errors.append("任务编号为空")
            elif task_id in ids:
                row_errors.append(f"任务编号重复：{task_id}")
            ids.add(task_id)
            video = _path(workbook.parent, video_s) if video_s else None
            if not video or not video.is_file():
                row_errors.append("视频路径不存在")
            elif video.suffix.lower() != ".mp4":
                row_errors.append("视频应为 MP4")
            elif video.stat().st_size > MAX_VIDEO_BYTES:
                row_errors.append("视频超过 1.5GB")
            if not title or len(title) > 30:
                row_errors.append("标题须为 1～30 字")
            if len(body) > 1000:
                row_errors.append("正文超过 1000 字")
            products = _split(products_s)
            if len(products) > 6 or len(set(products)) != len(products):
                row_errors.append("商品 ID 最多 6 个且不能重复")
            cover = _path(workbook.parent, cover_s) if cover_s else None
            if cover and (not cover.is_file() or cover.suffix.lower() not in {".jpg", ".jpeg", ".png"}):
                row_errors.append("封面路径须指向 JPG 或 PNG 文件")
            publish_at = None
            if time_s:
                try:
                    raw_time = raw[9]
                    parsed = raw_time if isinstance(raw_time, datetime) else datetime.strptime(time_s, TIME_FORMAT)
                    publish_at = parsed.replace(tzinfo=TZ)
                    if publish_at <= current:
                        row_errors.append("发布时间必须晚于当前北京时间")
                    elif publish_at < current + SCHEDULE_MIN_LEAD:
                        row_errors.append("定时发布至少应提前 2 小时")
                    elif publish_at > current + SCHEDULE_MAX_AHEAD:
                        row_errors.append("定时发布最多可设置未来 30 天")
                except ValueError:
                    row_errors.append("发布时间格式应为 YYYY-MM-DD HH:MM")
            if declaration not in DECLARATIONS:
                row_errors.append("创作者声明必须填写模板列出的选项之一")
            if not upload_platform:
                row_errors.append("上传平台不能为空")
            if row_errors:
                errors.append(f"第 {row_number} 行：" + "；".join(row_errors))
                continue
            assert video is not None
            tasks.append(Task(row_number, task_id, video, title, body, products, cover,
                              _split(brands_s), _split(tags_s), topic, publish_at, declaration,
                              upload_platform))
        if not tasks and not errors:
            errors.append("Excel 没有任务行")
        if errors:
            raise TaskValidationError(errors)
        return tasks
    finally:
        wb.close()
