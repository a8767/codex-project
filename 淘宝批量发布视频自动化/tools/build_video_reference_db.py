"""Combine authorized video and search exports into a traceable reference database.

Run from the project root: python tools/build_video_reference_db.py
The script reads source files only and replaces its own outputs in outputs/video_reference_db.
"""

from __future__ import annotations

import collections
import hashlib
import json
import re
import sqlite3
from pathlib import Path
import sys

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from taobao_uploader.copy_library import seed_copy_templates
from taobao_uploader.bike_copy_library import seed_bike_video_copies
from import_copy_bank import import_copy_bank, JSON_PATH as COPY_BANK_SOURCE
from import_guanghe_hot import import_capture, SOURCE as GUANGHE_HOT_SOURCE
SOURCE = ROOT / ".test_run"
OUT = ROOT / "outputs" / "video_reference_db"
DOWNLOADS = Path.home() / "Downloads"
STORE_FILE = DOWNLOADS / "新建任务09_24+13_50_00.xlsx"
HOT_FILES = sorted(
    p for p in DOWNLOADS.glob("行业热词明细数据_2026_08_24*.xlsx")
    if not p.name.startswith("~$")
)
TOPIC_RE = re.compile(r"[#＃]\s*([^\s#＃，,。！？!?；;、：:\[\]（）()]{1,60})")


def load_json(name: str):
    return json.loads((SOURCE / name).read_text(encoding="utf-8"))


def norm(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def find_topics(copy: str) -> tuple[str, list[str]]:
    topics = list(dict.fromkeys(norm(m.group(1)) for m in TOPIC_RE.finditer(copy)))
    body = norm(TOPIC_RE.sub(" ", copy))
    return body, [t for t in topics if t]


def no_id_key(row: dict) -> str:
    # The after-search chart exposes synthetic table keys, not actual video IDs.
    stable = "|".join((norm(row.get("copy")), norm(row.get("creator_id")), norm(row.get("published_at"))))
    return "douyin:unresolved:" + hashlib.sha256(stable.encode("utf-8")).hexdigest()[:24]


def sql_text(value) -> str:
    return "" if value is None else str(value)


def main() -> None:
    if not STORE_FILE.exists() or len(HOT_FILES) != 4:
        raise FileNotFoundError(f"Missing source: store={STORE_FILE.exists()}, hot files={len(HOT_FILES)}")
    OUT.mkdir(parents=True, exist_ok=True)
    db_path = OUT / "视频发布参考数据库.sqlite3"
    if db_path.exists():
        db_path.unlink()
    db = sqlite3.connect(db_path)
    db.executescript("""
        PRAGMA foreign_keys=ON;
        CREATE TABLE videos (
            video_key TEXT PRIMARY KEY, platform TEXT NOT NULL, video_id TEXT,
            title TEXT NOT NULL, body TEXT, creator TEXT, creator_id TEXT,
            published_at TEXT, video_url TEXT, shop_name TEXT, account_type TEXT,
            duration_seconds TEXT, availability TEXT NOT NULL DEFAULT '未验证',
            relevance_status TEXT NOT NULL DEFAULT '未筛选'
        );
        CREATE UNIQUE INDEX video_platform_id ON videos(platform,video_id) WHERE video_id IS NOT NULL;
        CREATE TABLE video_sources (
            video_key TEXT NOT NULL REFERENCES videos(video_key), source_type TEXT NOT NULL,
            source_file TEXT NOT NULL, source_row INTEGER, PRIMARY KEY(video_key,source_type,source_file,source_row)
        );
        CREATE TABLE video_topics (
            video_key TEXT NOT NULL REFERENCES videos(video_key), topic TEXT NOT NULL,
            extraction TEXT NOT NULL, PRIMARY KEY(video_key,topic)
        );
        CREATE TABLE rank_occurrences (
            video_key TEXT NOT NULL REFERENCES videos(video_key), chart TEXT NOT NULL,
            month INTEGER NOT NULL, rank INTEGER NOT NULL, video_type TEXT,
            metrics_json TEXT, PRIMARY KEY(chart,month,rank)
        );
        CREATE TABLE video_search_terms (
            video_key TEXT NOT NULL REFERENCES videos(video_key), chart TEXT NOT NULL,
            month INTEGER NOT NULL, rank INTEGER NOT NULL, display_order INTEGER NOT NULL,
            term TEXT NOT NULL, scope TEXT NOT NULL,
            PRIMARY KEY(chart,month,rank,display_order)
        );
        CREATE TABLE platform_search_term_records (
            source_file TEXT NOT NULL, source_row INTEGER NOT NULL, term TEXT NOT NULL,
            search_result_exposure TEXT, search_users TEXT, search_payment TEXT,
            product_exposure TEXT, product_ctr TEXT, product_conversion TEXT,
            PRIMARY KEY(source_file,source_row)
        );
        CREATE TABLE platform_search_terms (
            term TEXT PRIMARY KEY, source_count INTEGER NOT NULL, source_files TEXT NOT NULL
        );
        CREATE TABLE store_video_metrics (
            source_row INTEGER PRIMARY KEY, video_key TEXT NOT NULL REFERENCES videos(video_key),
            statistics_period TEXT, product_name TEXT, product_id TEXT, orders TEXT,
            payment TEXT, views TEXT, likes TEXT, comments TEXT, followers TEXT,
            raw_row_json TEXT NOT NULL
        );
        CREATE TABLE script_checks (
            chart TEXT NOT NULL, month INTEGER NOT NULL, rank INTEGER NOT NULL,
            video_key TEXT NOT NULL REFERENCES videos(video_key), status TEXT NOT NULL,
            transcript TEXT, evidence TEXT,
            PRIMARY KEY(chart,month,rank)
        );
    """)

    def put_video(key: str, platform: str, video_id, title, body="", creator="", creator_id="",
                  published_at="", video_url="", shop_name="", account_type="", duration_seconds=""):
        fields = (key, platform, sql_text(video_id) if video_id else None, norm(title), norm(body),
                  norm(creator), norm(creator_id), norm(published_at), norm(video_url),
                  norm(shop_name), norm(account_type), sql_text(duration_seconds))
        db.execute("""INSERT INTO videos(video_key,platform,video_id,title,body,creator,creator_id,
                   published_at,video_url,shop_name,account_type,duration_seconds)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(video_key) DO UPDATE SET
                   title=CASE WHEN length(excluded.title)>length(videos.title) THEN excluded.title ELSE videos.title END,
                   body=CASE WHEN length(excluded.body)>length(videos.body) THEN excluded.body ELSE videos.body END,
                   creator=CASE WHEN videos.creator='' THEN excluded.creator ELSE videos.creator END,
                   creator_id=CASE WHEN videos.creator_id='' THEN excluded.creator_id ELSE videos.creator_id END,
                   published_at=CASE WHEN length(excluded.published_at)>length(videos.published_at)
                     THEN excluded.published_at ELSE videos.published_at END,
                   video_url=CASE WHEN videos.video_url='' THEN excluded.video_url ELSE videos.video_url END,
                   shop_name=CASE WHEN videos.shop_name='' THEN excluded.shop_name ELSE videos.shop_name END,
                   account_type=CASE WHEN videos.account_type='' THEN excluded.account_type ELSE videos.account_type END,
                   duration_seconds=CASE WHEN videos.duration_seconds='' THEN excluded.duration_seconds ELSE videos.duration_seconds END""", fields)

    def add_topics(key: str, topics: list[str], extraction: str):
        for topic in topics:
            if topic:
                db.execute("INSERT OR IGNORE INTO video_topics VALUES(?,?,?)", (key, topic, extraction))

    def source(key: str, kind: str, filename: str, row: int):
        db.execute("INSERT OR IGNORE INTO video_sources VALUES(?,?,?,?)", (key, kind, filename, row))

    for i, row in enumerate(load_json("quality_2026_rows.json"), 1):
        key = "taobao:" + str(row["id"])
        put_video(key, "淘宝光合", row["id"], row["title"], row["title"], published_at=row["published_at"])
        add_topics(key, row.get("topics", []), "平台话题")
        source(key, "淘宝优质视频", "quality_2026_rows.json", i)

    store_book = load_workbook(STORE_FILE, read_only=False, data_only=True)
    store_sheet = store_book.active
    store_headers = [sql_text(c.value) for c in store_sheet[1]]
    store_count = 0
    for rowno, values in enumerate(store_sheet.iter_rows(min_row=2, values_only=True), 2):
        if not values[1]:
            continue
        store_count += 1
        vid = str(values[1]).strip()
        key = "douyin:" + vid
        title = sql_text(values[2] or values[10])
        body, topics = find_topics(title)
        put_video(key, "抖音", vid, title, body, values[7], values[5], values[11],
                  values[13], values[3], values[8], values[12])
        add_topics(key, topics, "标题#话题")
        source(key, "店铺视频导出", STORE_FILE.name, rowno)
        raw = {f"{h}_{idx}" if h == "视频标题" else h: sql_text(v)
               for idx, (h, v) in enumerate(zip(store_headers, values))}
        db.execute("INSERT INTO store_video_metrics VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                   (rowno, key, sql_text(values[0]), sql_text(values[14]), sql_text(values[15]),
                    sql_text(values[17]), sql_text(values[18]), sql_text(values[20]),
                    sql_text(values[27]), sql_text(values[28]), sql_text(values[29]),
                    json.dumps(raw, ensure_ascii=False)))
    store_book.close()

    sales = load_json("douyin_top100_2026_07_09.json")
    new_charts = load_json("douyin_hot_search_2026_07_09.json")
    skipped_invalid = 0
    # Exact copy + creator + displayed publish time is used only when it identifies one real ID.
    known = collections.defaultdict(set)
    for row in sales + new_charts:
        if row.get("video_id") and norm(row.get("copy")) != "视频已失效":
            known[(norm(row.get("copy")), norm(row.get("creator_id")),
                   norm(row.get("published_at")))].add(str(row["video_id"]))
    for i, row in enumerate(sales + new_charts, 1):
        if norm(row.get("copy")) == "视频已失效":
            skipped_invalid += 1
            continue
        vid = str(row["video_id"]) if row.get("video_id") else None
        if not vid:
            matches = known.get((norm(row.get("copy")), norm(row.get("creator_id")),
                                 norm(row.get("published_at"))), set())
            if len(matches) == 1:
                vid = next(iter(matches))
        key = "douyin:" + vid if vid else no_id_key(row)
        copy = norm(row.get("copy"))
        body, topics = find_topics(copy)
        put_video(key, "抖音", vid, copy, body, row.get("creator"),
                  row.get("creator_id"), row.get("published_at"),
                  f"https://www.douyin.com/video/{vid}" if vid else "")
        add_topics(key, topics, "榜单文案#话题")
        chart = row.get("chart", "视频销量榜")
        month, rank = int(row["month"]), int(row["rank"])
        source(key, chart, "douyin_top100_2026_07_09.json" if i <= len(sales)
               else "douyin_hot_search_2026_07_09.json", i if i <= len(sales) else i-len(sales))
        metrics = row.get("metrics") or {k: row.get(k, "") for k in
                 ("views", "payment", "likes", "comments", "shares")}
        db.execute("INSERT INTO rank_occurrences VALUES(?,?,?,?,?,?)",
                   (key, chart, month, rank, row.get("video_type", "自营"),
                    json.dumps(metrics, ensure_ascii=False)))
        for pos, term in enumerate(row.get("visible_search_terms", []), 1):
            term = norm(term)
            if term:
                db.execute("INSERT INTO video_search_terms VALUES(?,?,?,?,?,?,?)",
                           (key, chart, month, rank, pos, term, "榜单页面展示的前两项"))
        if chart in ("看后搜视频榜", "热门视频榜") and rank <= 10:
            db.execute("INSERT INTO script_checks VALUES(?,?,?,?,?,?,?)",
                       (chart, month, rank, key, "待验证视频内容", "", "浏览器视频页要求登录；不能从标题推断脚本"))

    term_files = collections.defaultdict(set)
    hot_raw_count = 0
    for path in HOT_FILES:
        wb = load_workbook(path, read_only=False, data_only=True)
        for rowno, row in enumerate(wb.active.iter_rows(min_row=2, values_only=True), 2):
            term = norm(row[0])
            if not term:
                continue
            hot_raw_count += 1
            values = [sql_text(v) for v in row[:7]]
            db.execute("INSERT INTO platform_search_term_records VALUES(?,?,?,?,?,?,?,?,?)",
                       (path.name, rowno, *values))
            term_files[term].add(path.name)
        wb.close()
    for term, files in term_files.items():
        db.execute("INSERT INTO platform_search_terms VALUES(?,?,?)",
                   (term, len(files), "；".join(sorted(files))))

    retest_path = SOURCE / "script_retest_2026_09_24.json"
    if retest_path.exists():
        for check in json.loads(retest_path.read_text(encoding="utf-8")):
            db.execute("""UPDATE script_checks SET status=?,transcript=?,evidence=?
                       WHERE video_key=?""",
                       (check["status"], check["transcript"], check["evidence"], check["video_key"]))

    copy_count = seed_copy_templates(db)
    bike_copy_count = seed_bike_video_copies(db)
    copy_bank_count = import_copy_bank(db) if COPY_BANK_SOURCE.exists() else 0
    guanghe_hot = import_capture(db) if GUANGHE_HOT_SOURCE.exists() else None
    db.commit()
    checks = {
        "taobao_quality_rows": 320,
        "store_export_rows": store_count,
        "sales_rank_rows": len(sales),
        "new_rank_rows": len(new_charts),
        "invalid_rank_rows_skipped": skipped_invalid,
        "valid_rank_rows": db.execute("SELECT COUNT(*) FROM rank_occurrences").fetchone()[0],
        "unique_videos": db.execute("SELECT COUNT(*) FROM videos").fetchone()[0],
        "video_topics": db.execute("SELECT COUNT(*) FROM video_topics").fetchone()[0],
        "visible_search_term_links": db.execute("SELECT COUNT(*) FROM video_search_terms").fetchone()[0],
        "platform_search_term_raw_rows": hot_raw_count,
        "platform_search_terms_unique": len(term_files),
        "script_targets": db.execute("SELECT COUNT(*) FROM script_checks").fetchone()[0],
        "unresolved_video_ids": db.execute("SELECT COUNT(*) FROM videos WHERE video_key LIKE 'douyin:unresolved:%'").fetchone()[0],
        "original_copy_templates": copy_count,
        "bike_video_copies": bike_copy_count,
        "original_copy_bank": copy_bank_count,
        "guanghe_hot": guanghe_hot,
    }
    assert checks["valid_rank_rows"] == 900 - skipped_invalid
    assert checks["script_targets"] <= 60
    (OUT / "数据校验.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"database": str(db_path), **checks}, ensure_ascii=False))
    db.close()


if __name__ == "__main__":
    main()
